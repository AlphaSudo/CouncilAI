"""Bounded, CPU-only document extraction in an isolated, monitored process."""

import base64
import io
import json
import os
import subprocess
import sys
import threading
import zipfile

LIMIT = 2 * 1024 * 1024
PARSER_SLOTS = threading.BoundedSemaphore(2)


def extract_inner(raw, mime):
    pages, limitations = [], []
    if mime == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF unsupported")
        for i, page in enumerate(reader.pages):
            if i >= 100:
                limitations.append("Only first 100 pages extracted")
                break
            text = page.extract_text() or ""
            pages.append(f"[Page {i + 1}]\n" + text)
            if sum(len(p) for p in pages) > LIMIT:
                break
        if not any(p.split("\n", 1)[-1].strip() for p in pages):
            limitations.append("No extractable text; OCR not available")
        parser = "pypdf/6.21.1"
    elif (
        mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if (
                sum(f.file_size for f in archive.infolist()) > 16 * LIMIT
                or len(archive.infolist()) > 1000
            ):
                raise ValueError("DOCX expansion bound exceeded")
        document = Document(io.BytesIO(raw))
        pages = [
            f"[Paragraph {i + 1}] {p.text}" for i, p in enumerate(document.paragraphs)
        ]
        for i, table in enumerate(document.tables):
            pages.append(
                f"[Table {i + 1}]\n"
                + "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            )
        limitations.append(
            "DOCX paragraphs/tables; headers, comments, shapes and visual pagination not extracted"
        )
        parser = "python-docx/1.2.0"
    elif mime in ("text/html", "application/xhtml+xml"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "iframe", "noscript", "svg", "form"]):
            tag.decompose()
        for heading in soup.select("h1,h2,h3,h4,h5,h6"):
            heading.replace_with(
                soup.new_string(
                    "\n"
                    + "#" * int(heading.name[1])
                    + " "
                    + heading.get_text(" ", strip=True)
                    + "\n"
                )
            )
        for row in soup.select("tr"):
            row.replace_with(
                soup.new_string(
                    "\n"
                    + " | ".join(
                        cell.get_text(" ", strip=True) for cell in row.select("th,td")
                    )
                    + "\n"
                )
            )
        pages = [soup.get_text("\n", strip=True)]
        limitations.append(
            "Static HTML text only; scripts and visual layout not executed"
        )
        parser = "beautifulsoup4/4.15.0/static-text-2"
    elif mime in ("text/plain", "text/markdown"):
        pages = [raw.decode("utf-8-sig", errors="replace")]
        parser = "utf8/1"
        if "\ufffd" in pages[0]:
            limitations.append("Invalid UTF-8 bytes replaced")
    else:
        raise ValueError("Unsupported source MIME type")
    text = "\n\n".join(pages)
    truncated = len(text.encode()) > LIMIT
    if truncated:
        text = text.encode()[:LIMIT].decode("utf-8", errors="ignore")
        limitations.append("Extracted text reached 2 MiB limit")
    return {
        "text": text,
        "parser": parser,
        "limitations": limitations,
        "coverage": "partial" if limitations else "text_complete",
        "truncated": truncated,
    }


def _extract(raw, mime):
    if len(raw) > LIMIT:
        raise ValueError("Source exceeds 2 MiB input limit")
    if mime in ("text/plain", "text/markdown"):
        return extract_inner(raw, mime)
    import psutil

    env = {**os.environ, "PYTHONUTF8": "1"}
    child = subprocess.Popen(
        [sys.executable, "-m", "zerowall.source_parsers"],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    stop = threading.Event()

    def monitor():
        while not stop.wait(0.02):
            try:
                if psutil.Process(child.pid).memory_info().rss > 384 * 1024 * 1024:
                    child.kill()
                    return
            except psutil.NoSuchProcess:
                return

    thread = threading.Thread(target=monitor, daemon=True)
    thread.start()
    try:
        output, _ = child.communicate(
            json.dumps({"raw": base64.b64encode(raw).decode(), "mime": mime}).encode(),
            timeout=10,
        )
        if child.returncode or len(output) > 4 * LIMIT:
            raise ValueError("Document parser failed or exceeded its resource limit")
        result = json.loads(output)
        if result.get("error"):
            raise ValueError(result["error"])
        return result
    except subprocess.TimeoutExpired:
        child.kill()
        child.communicate()
        raise ValueError("Document extraction deadline exceeded") from None
    finally:
        stop.set()
        thread.join(timeout=1)


def extract(raw, mime):
    if mime in ("text/plain", "text/markdown"):
        return _extract(raw, mime)
    if not PARSER_SLOTS.acquire(timeout=10):
        raise ValueError("Document parser capacity is busy")
    try:
        return _extract(raw, mime)
    finally:
        PARSER_SLOTS.release()


if __name__ == "__main__":
    try:
        request = json.loads(sys.stdin.buffer.read(4 * LIMIT))
        result = extract_inner(
            base64.b64decode(request["raw"], validate=True), request["mime"]
        )
    except Exception as error:  # noqa: BLE001 -- isolated parser boundary returns no document/exception contents.
        result = {"error": type(error).__name__ + ": source extraction failed"}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode())
